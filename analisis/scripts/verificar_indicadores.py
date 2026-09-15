from __future__ import annotations

import csv
import math
from pathlib import Path

from openpyxl import load_workbook


RAIZ = Path(__file__).resolve().parents[2]
LIBRO = RAIZ / "datos" / "procesados" / "Datos_CIMAC_procesados.xlsx"
SALIDA = RAIZ / "analisis" / "resultados" / "verificacion_2024.csv"


def cercano(obtenido: float, esperado: float, tolerancia: float = 5e-7) -> bool:
    return math.isclose(obtenido, esperado, rel_tol=0.0, abs_tol=tolerancia)


def main() -> None:
    libro = load_workbook(LIBRO, data_only=True, read_only=True)
    actual = libro["Actualizado"]
    resinas = libro["Resinas"]
    emisiones = libro["Emisiones"]
    metas = libro["Metas_REP"]

    resultados = {
        "TR_total": actual["N4"].value,
        "TR_packaging": actual["N9"].value,
        "Tgr_kg_hab_ano": actual["N7"].value * 1000 / actual["N22"].value,
        "IC_REP_est": actual["N11"].value / (metas["C3"].value * actual["M7"].value),
        "eC_PET": resinas["K3"].value / resinas["K21"].value,
        "eC_PP": resinas["K4"].value / resinas["K22"].value,
        "eC_PE": resinas["K5"].value / resinas["K26"].value,
    }

    irp_componentes = {}
    for fila, resina in ((3, "PET"), (4, "PP"), (5, "PE")):
        cantidad = resinas.cell(fila, 11).value
        diferencia_factor = emisiones.cell(fila, 2).value - emisiones.cell(fila, 3).value
        irp_componentes[resina] = diferencia_factor * cantidad / 1_000_000_000
    resultados["IRP_total_Mt_CO2eq"] = sum(irp_componentes.values())

    esperados = {
        "TR_total": 0.103,
        "TR_packaging": 0.130,
        "Tgr_kg_hab_ano": 30.065615230249822,
        "IC_REP_est": 0.819658772599949,
        "eC_PET": 0.21926115737774926,
        "eC_PP": 0.3193125,
        "eC_PE": 0.16852433281004708,
        "IRP_total_Mt_CO2eq": 0.174775335,
    }

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    filas = []
    fallos = []
    for indicador, esperado in esperados.items():
        obtenido = resultados[indicador]
        estado = "OK" if cercano(obtenido, esperado) else "DIFERENCIA"
        filas.append((indicador, obtenido, esperado, obtenido - esperado, estado))
        if estado != "OK":
            fallos.append(indicador)

    with SALIDA.open("w", newline="", encoding="utf-8") as archivo:
        escritor = csv.writer(archivo)
        escritor.writerow(["indicador", "obtenido", "esperado", "diferencia", "estado"])
        escritor.writerows(filas)

    if fallos:
        raise SystemExit("Falló la verificación: " + ", ".join(fallos))
    print(f"Verificación correcta: {len(filas)} resultados. Salida: {SALIDA}")


if __name__ == "__main__":
    main()
